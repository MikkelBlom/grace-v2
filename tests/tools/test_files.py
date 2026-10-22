"""The files group (services/grace_tools/files.py, files_docs.py): every file lives under tmp_path,
home is a fake folder, the Recycle Bin is a folder, and Office/LibreOffice are switched off so the
built-in converters are what is tested."""

import json
import os
import time
import zipfile
from pathlib import Path

import pytest

import grace_tools
from grace_tools import registry
from grace_tools import files as F
from grace_tools import files_docs as D


def call(ctx, tool_name, confirmed=False, **args):
    """conftest.call, but these tools take an argument called `name` (the file name), which
    conftest.call uses for the tool's own name."""
    out = registry.run(ctx, tool_name, args, lambda: True, confirmed)
    assert out is not None, f"no tool {tool_name}"
    assert isinstance(out.get("result"), str), out
    return out


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch):
    grace_tools.load()
    h = tmp_path / "home"
    for r in F.ROOT_NAMES:
        (h / r).mkdir(parents=True)
    monkeypatch.setattr(F, "HOME", h)
    trash = tmp_path / "trash"
    trash.mkdir()

    def to_trash(p):
        p = Path(p)
        p.rename(trash / f"{len(list(trash.iterdir()))}-{p.name}")
    monkeypatch.setattr(F, "send_to_trash", to_trash)
    monkeypatch.setattr(D, "office_installed", lambda app: False)
    monkeypatch.setattr(D, "find_soffice", lambda: None)
    D.CACHE.clear()
    return h


def grace_dir(home):
    return home / "Documents" / "Grace"


# ---- paths and safety -------------------------------------------------------------------------

def test_write_file_goes_to_documents_grace(ctx, home):
    out = call(ctx, "write_file", name="indkøbsliste", content="# Indkøb\n- mælk\n- æg\n- rugbrød")
    p = grace_dir(home) / "indkøbsliste.md"
    assert out["ok"] and p.read_text("utf-8").startswith("# Indkøb")
    assert str(p) in out["result"] and out["risk"] == "visible"


def test_write_file_plain_text_and_danish_folder(ctx, home):
    out = call(ctx, "write_file", name="huskeseddel", content="Ring til tandlægen", folder="skrivebordet")
    assert out["ok"] and (home / "Desktop" / "huskeseddel.txt").exists()
    out = call(ctx, "write_file", name="dokumenter/noter.md", content="hej")
    assert out["ok"] and (home / "Documents" / "noter.md").exists()


def test_write_file_never_overwrites(ctx, home):
    call(ctx, "write_file", name="a.txt", content="first")
    out = call(ctx, "write_file", name="a.txt", content="second")
    assert not out["ok"] and "already exists" in out["result"]
    assert (grace_dir(home) / "a.txt").read_text("utf-8") == "first"


def test_write_file_dry_run_changes_nothing(dry_ctx, home):
    out = call(dry_ctx, "write_file", name="x.txt", content="hej")
    assert out["ok"] and out["result"].startswith("Would")
    assert not grace_dir(home).exists()


@pytest.mark.parametrize("name", [r"C:\Windows\System32\evil.txt", "~/AppData/Roaming/x.txt", "~/.ssh/config.txt",
                                  "script.bat", "~/Documents/run.ps1"])
def test_refuses_outside_home_system_folders_and_scripts(ctx, home, name):
    out = call(ctx, "write_file", name=name, content="x")
    assert not out["ok"] and "Nothing was changed" in out["result"]


def test_bare_name_is_found_in_grace_folder(ctx, home):
    call(ctx, "write_file", name="noter.txt", content="linje")
    out = call(ctx, "read_document", path="noter.txt")
    assert out["ok"] and "linje" in out["result"]


def test_bare_name_falls_back_to_file_index(ctx, home):
    p = home / "Pictures" / "ferie" / "plan.txt"
    p.parent.mkdir()
    p.write_text("Rom i maj", "utf-8")

    class Index:
        def candidates(self, description, limit=20, kind="any"):
            return [(5, str(p), False, 0.0)] if "plan" in description else []
    ctx.files = Index()
    assert "Rom i maj" in call(ctx, "read_document", path="ferieplan")["result"]
    assert not call(ctx, "read_document", path="noget andet")["ok"]


# ---- append, edit, undo -----------------------------------------------------------------------

def test_append_and_undo(ctx, home):
    call(ctx, "write_file", name="liste.txt", content="mælk")
    out = call(ctx, "append_to_file", path="liste.txt", text="smør")
    p = grace_dir(home) / "liste.txt"
    assert out["ok"] and p.read_text("utf-8") == "mælk\nsmør\n"
    assert call(ctx, "undo_last")["ok"]
    assert p.read_text("utf-8") == "mælk"
    assert call(ctx, "undo_last")["ok"] and not p.exists()  # then the write itself
    assert len(list((home.parent / "trash").iterdir())) == 1  # undoing a write uses the Recycle Bin too
    assert not call(ctx, "undo_last")["ok"]


def test_append_to_word(ctx, home):
    call(ctx, "create_document", format="docx", title="Indkøb", sections=[{"bullets": ["mælk"]}])
    p = grace_dir(home) / "Indkøb.docx"
    assert call(ctx, "append_to_file", path=str(p), text="- kaffe\nHusk posen")["ok"]
    text = call(ctx, "read_document", path=str(p))["result"]
    assert "- mælk" in text and "- kaffe" in text and "Husk posen" in text


def test_edit_asks_first_then_replaces(ctx, home):
    call(ctx, "write_file", name="brev.txt", content="Kære Peter\nVi ses fredag.\n")
    p = grace_dir(home) / "brev.txt"
    out = call(ctx, "edit_file", path=str(p), find="fredag", replace="lørdag")
    assert out["needs_confirm"] and "fredag" in out["confirm"]["summary"] and "fredag" in p.read_text("utf-8")
    out = call(ctx, "edit_file", confirmed=True, path=str(p), find="FREDAG", replace="lørdag")  # STT casing
    assert out["ok"] and p.read_text("utf-8") == "Kære Peter\nVi ses lørdag.\n"
    assert call(ctx, "undo_last")["ok"] and "fredag" in p.read_text("utf-8")


def test_edit_not_found_and_whole_content(ctx, home):
    call(ctx, "write_file", name="n.md", content="gammel")
    p = grace_dir(home) / "n.md"
    out = call(ctx, "edit_file", confirmed=True, path=str(p), find="findes ikke", replace="x")
    assert not out["ok"] and "isn't in" in out["result"]
    assert call(ctx, "edit_file", confirmed=True, path=str(p), new_content="ny tekst")["ok"]
    assert p.read_text("utf-8") == "ny tekst"


def test_edit_word_document(ctx, home):
    call(ctx, "create_document", format="docx", title="Plan", sections=[{"heading": "Mandag", "paragraphs": ["Møde kl 10"]}])
    p = grace_dir(home) / "Plan.docx"
    assert call(ctx, "edit_file", confirmed=True, path=str(p), find="kl 10", replace="kl 11")["ok"]
    assert "Møde kl 11" in call(ctx, "read_document", path=str(p))["result"]


def test_edit_dry_run(dry_ctx, home):
    (home / "Documents" / "a.txt").write_text("abc", "utf-8")
    out = call(dry_ctx, "edit_file", confirmed=True, path="dokumenter/a.txt", find="b", replace="x")
    assert out["result"].startswith("Would") and (home / "Documents" / "a.txt").read_text("utf-8") == "abc"


def test_undo_log_and_backups_stay_small(ctx, home, monkeypatch):
    monkeypatch.setattr(F, "UNDO_KEEP", 3)
    call(ctx, "write_file", name="b.txt", content="0")
    for i in range(5):
        call(ctx, "append_to_file", path="b.txt", text=str(i))
    assert len((ctx.data_dir / "files-undo.jsonl").read_text("utf-8").splitlines()) == 3
    assert len(list((ctx.data_dir / "files-undo").iterdir())) == 3


def test_undo_refuses_after_his_own_change(ctx, home):
    call(ctx, "write_file", name="u.txt", content="a")
    call(ctx, "append_to_file", path="u.txt", text="b")
    p = grace_dir(home) / "u.txt"
    time.sleep(0.02)
    p.write_text("his own work", "utf-8")
    out = call(ctx, "undo_last")
    assert not out["ok"] and "changed after" in out["result"] and p.read_text("utf-8") == "his own work"


# ---- move, rename, copy, replace, folders ------------------------------------------------------

def test_rename_keeps_extension_and_undo(ctx, home):
    src = home / "Documents" / "rapport.docx"
    src.write_bytes(b"x")
    out = call(ctx, "rename_file", path=str(src), new_name="rapport endelig")
    assert out["ok"] and (home / "Documents" / "rapport endelig.docx").exists() and not src.exists()
    assert call(ctx, "undo_last")["ok"] and src.exists()


def test_rename_refuses_script_extension(ctx, home):
    src = home / "Documents" / "a.txt"
    src.write_text("x")
    assert not call(ctx, "rename_file", path=str(src), new_name="a.bat")["ok"] and src.exists()


def test_move_to_danish_folder_and_refuse_existing(ctx, home):
    src = home / "Downloads" / "faktura.pdf"
    src.write_bytes(b"1")
    out = call(ctx, "move_file", path=str(src), to="skrivebordet")
    assert out["ok"] and (home / "Desktop" / "faktura.pdf").exists()
    again = home / "Downloads" / "faktura.pdf"
    again.write_bytes(b"2")
    out = call(ctx, "move_file", path=str(again), to="skrivebord")
    assert not out["ok"] and "replace_file" in out["result"] and again.exists()
    assert (home / "Desktop" / "faktura.pdf").read_bytes() == b"1"


def test_move_into_missing_folder_says_make_it(ctx, home):
    src = home / "Downloads" / "a.txt"
    src.write_text("x")
    out = call(ctx, "move_file", path=str(src), to="dokumenter/Projekter")
    assert not out["ok"] and "make_folder" in out["result"]


def test_main_folders_cannot_be_moved(ctx, home):
    out = call(ctx, "move_file", path="skrivebordet", to="dokumenter")
    assert not out["ok"] and "main folders" in out["result"]


def test_copy_and_replace(ctx, home):
    a, b = home / "Documents" / "a.txt", home / "Desktop" / "a.txt"
    a.write_text("ny", "utf-8")
    b.write_text("gammel", "utf-8")
    out = call(ctx, "copy_file", path=str(a), to="skrivebordet")
    assert not out["ok"] and b.read_text("utf-8") == "gammel"
    out = call(ctx, "replace_file", path=str(a), target=str(b))
    assert out["needs_confirm"] and "replacing" in out["confirm"]["summary"]
    out = call(ctx, "replace_file", confirmed=True, path=str(a), target=str(b))
    assert out["ok"] and b.read_text("utf-8") == "ny" and not a.exists()
    assert call(ctx, "undo_last")["ok"]
    assert a.read_text("utf-8") == "ny" and b.read_text("utf-8") == "gammel"


def test_copy_with_new_name(ctx, home):
    a = home / "Documents" / "cv.docx"
    a.write_bytes(b"cv")
    assert call(ctx, "copy_file", path=str(a), to="cv kopi.docx")["ok"]
    assert (home / "Documents" / "cv kopi.docx").read_bytes() == b"cv" and a.exists()


def test_make_folder_and_undo(ctx, home):
    out = call(ctx, "make_folder", name="Projekter/2026", folder="dokumenter")
    p = home / "Documents" / "Projekter" / "2026"
    assert out["ok"] and p.is_dir()
    assert call(ctx, "undo_last")["ok"] and not (home / "Documents" / "Projekter").exists()


# ---- delete -----------------------------------------------------------------------------------

def test_delete_asks_first_then_uses_recycle_bin(ctx, home):
    files = [home / "Downloads" / f"{n}.pdf" for n in ("a", "b", "c")]
    for f in files:
        f.write_bytes(b"x")
    out = call(ctx, "delete_files", paths=[str(f) for f in files])
    assert out["needs_confirm"] and "3 items" in out["confirm"]["summary"] and "Recycle Bin" in out["confirm"]["summary"]
    assert all(f.exists() for f in files)
    out = call(ctx, "delete_files", confirmed=True, paths=[str(f) for f in files])
    assert out["ok"] and not any(f.exists() for f in files) and len(list((home.parent / "trash").iterdir())) == 3
    assert "Recycle Bin" in call(ctx, "undo_last")["result"]


def test_delete_accepts_a_single_string(ctx, home):
    f = home / "Desktop" / "gammel.txt"
    f.write_text("x")
    assert call(ctx, "delete_files", confirmed=True, paths=str(f))["ok"] and not f.exists()


@pytest.mark.parametrize("target", ["dokumenter", "~", r"C:\Windows\notepad.exe", "~/AppData"])
def test_delete_refuses_main_and_system_folders(ctx, home, target):
    (home / "AppData").mkdir(exist_ok=True)
    out = call(ctx, "delete_files", confirmed=True, paths=[target])
    assert not out["ok"] and not list((home.parent / "trash").iterdir())


def test_delete_dry_run(dry_ctx, home):
    f = home / "Desktop" / "x.txt"
    f.write_text("x")
    out = call(dry_ctx, "delete_files", confirmed=True, paths=[str(f)])
    assert out["result"].startswith("Would") and f.exists()


# ---- zip --------------------------------------------------------------------------------------

def test_zip_and_unzip(ctx, home):
    d = home / "Documents" / "Opgave"
    (d / "bilag").mkdir(parents=True)
    (d / "opgave.md").write_text("# Opgave", "utf-8")
    (d / "bilag" / "data.csv").write_text("a,b\n1,2", "utf-8")
    out = call(ctx, "zip_files", paths=[str(d)])
    z = home / "Documents" / "Opgave.zip"
    assert out["ok"] and z.exists() and "2 file(s)" in out["result"]
    out = call(ctx, "unzip_file", path=str(z), to=str(home / "Desktop" / "udpakket"))
    assert out["ok"] and (home / "Desktop" / "udpakket" / "Opgave" / "bilag" / "data.csv").exists()


def test_unzip_refuses_zip_slip(ctx, home):
    z = home / "Downloads" / "ond.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("../../escape.txt", "x")
    out = call(ctx, "unzip_file", path=str(z))
    assert not out["ok"] and not (home.parent / "escape.txt").exists()


# ---- read_document ------------------------------------------------------------------------------

def test_read_long_text_in_parts(ctx, home):
    p = home / "Documents" / "lang.txt"
    p.write_text("\n".join(f"Linje {i} med lidt tekst om ingenting" for i in range(300)), "utf-8")
    out = call(ctx, "read_document", path=str(p))
    assert out["detail"]["parts"] > 1 and "part=2" in out["result"] and len(out["result"]) <= 1500
    last = call(ctx, "read_document", path=str(p), part=out["detail"]["parts"])
    assert "Linje 299" in last["result"] and "End of the file" in last["result"]
    assert not call(ctx, "read_document", path=str(p), part=999)["ok"]


def test_read_spreadsheets_slides_csv_json(ctx, home):
    call(ctx, "create_document", format="xlsx", title="Budget", table=[["Post", "Beløb"], ["Husleje", "8.500"], ["Mad", "3000"]])
    r = call(ctx, "read_document", path="Budget.xlsx")["result"]
    assert "Husleje | 8500" in r and "Excel" in r
    call(ctx, "create_document", format="pptx", title="Oplæg", slides=[{"title": "Mål", "bullets": ["Hurtig", "Lokal"]}])
    r = call(ctx, "read_document", path="Oplæg.pptx")["result"]
    assert "Slide 2: Mål" in r and "Lokal" in r
    (home / "Documents" / "d.csv").write_text("navn;alder\nSofie;30\n", "utf-8")
    assert "Sofie | 30" in call(ctx, "read_document", path="dokumenter/d.csv")["result"]
    (home / "Documents" / "d.json").write_text(json.dumps({"by": "Aarhus"}), "utf-8")
    assert "Aarhus" in call(ctx, "read_document", path="dokumenter/d.json")["result"]


def test_read_pdf_and_scanned_pdf(ctx, home):
    call(ctx, "create_document", format="pdf", title="Kontrakt", sections=[{"heading": "Pris", "paragraphs": ["Prisen er 1.200 kr"]}])
    assert "Prisen er 1.200 kr" in call(ctx, "read_document", path="Kontrakt.pdf")["result"]
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.output(str(home / "Documents" / "scan.pdf"))
    out = call(ctx, "read_document", path="dokumenter/scan.pdf")
    assert out["ok"] and "no text layer" in out["result"]


def test_read_says_old_formats_and_pictures(ctx, home):
    (home / "Documents" / "gammel.doc").write_bytes(b"\xd0\xcf")
    assert "old Office format" in call(ctx, "read_document", path="dokumenter/gammel.doc")["result"]
    from PIL import Image
    Image.new("RGB", (40, 30)).save(home / "Pictures" / "a.png")
    assert "40 x 30" in call(ctx, "read_document", path="billeder/a.png")["result"]


# ---- create_document and charts -----------------------------------------------------------------

def test_create_word_with_everything(ctx, home):
    out = call(ctx, "create_document", format="word", title="Indkøbsliste",
               sections=[{"heading": "Mejeri", "bullets": ["Mælk", "Ost"]},
                         {"heading": "Priser", "table": [["Vare", "Pris"], ["Mælk", "12,5"]],
                          "chart": {"kind": "bar", "labels": ["Mælk", "Ost"], "values": [12.5, 40]}}])
    p = grace_dir(home) / "Indkøbsliste.docx"
    assert out["ok"] and p.exists() and "1 chart" in out["result"]
    import docx
    d = docx.Document(str(p))
    assert len(d.inline_shapes) == 1 and d.tables[0].cell(1, 1).text == "12,5"
    assert not list(ctx.scratch_dir.glob("chart-*"))  # temporary chart files are cleaned up


def test_create_accepts_json_strings_and_never_overwrites(ctx, home):
    args = dict(format="md", title="Noter", sections=json.dumps([{"heading": "A", "paragraphs": "første\n\nanden"}]))
    assert call(ctx, "create_document", **args)["ok"]
    out = call(ctx, "create_document", **args)
    assert out["ok"] and "Noter (2).md" in out["result"]
    text = (grace_dir(home) / "Noter.md").read_text("utf-8")
    assert "## A" in text and "første" in text and "anden" in text


def test_create_pdf_keeps_danish_letters(ctx, home):
    call(ctx, "create_document", format="pdf", title="Æbler og rødgrød", sections=[{"bullets": ["blåbær – 20 kr €"]}])
    r = call(ctx, "read_document", path="Æbler og rødgrød.pdf")["result"]
    assert "blåbær" in r and "Æbler" in r


def test_create_excel_numbers_and_chart(ctx, home):
    call(ctx, "create_document", format="xlsx", title="Træning", table=[["Uge", "Km"], ["1", "12,5"], ["2", "15"]],
         chart={"kind": "line", "labels": ["1", "2"], "values": ["12,5", "15"]})
    import openpyxl
    ws = openpyxl.load_workbook(grace_dir(home) / "Træning.xlsx").active
    assert ws["B2"].value == 12.5 and ws["B3"].value == 15 and len(ws._images) == 1


def test_create_powerpoint_from_sections(ctx, home):
    out = call(ctx, "create_document", format="pptx", title="Ferie",
               sections=[{"heading": "Hvor", "bullets": ["Rom"]}, {"heading": "Budget", "table": [["Post", "Kr"], ["Fly", "2000"]]}])
    assert out["ok"] and "2 slides" in out["result"]


def test_create_needs_content_and_dry_run(ctx, dry_ctx, home):
    assert not call(ctx, "create_document", format="docx", title="Tom")["ok"]
    out = call(dry_ctx, "create_document", format="docx", title="X", sections=[{"paragraphs": ["a"]}])
    assert out["result"].startswith("Would") and not grace_dir(home).exists()


def test_make_chart_pie_folds_into_other(ctx, home):
    labels = [f"P{i}" for i in range(9)]
    out = call(ctx, "make_chart", kind="pie", title="Udgifter", labels=labels, values=list(range(1, 10)))
    p = grace_dir(home) / "Udgifter.png"
    assert out["ok"] and p.exists() and p.read_bytes()[:4] == b"\x89PNG"


def test_make_chart_bad_values(ctx, home):
    out = call(ctx, "make_chart", kind="bar", labels=["a", "b"], values=["mange", 2])
    assert not out["ok"] and "numbers" in out["result"]
    out = call(ctx, "make_chart", kind="bar", labels=["a", "b", "c"], values=[1, 2])
    assert not out["ok"] and "labels" in out["result"]


def test_make_chart_several_series(ctx, home):
    out = call(ctx, "make_chart", kind="bar", title="Salg", labels=["jan", "feb"],
               series=[{"name": "2025", "values": [1, 2]}, {"name": "2026", "values": [2, 3]}])
    assert out["ok"] and "2 series" in out["result"]


# ---- compare --------------------------------------------------------------------------------------

def test_compare_text(ctx, home):
    a, b = home / "Documents" / "v1.txt", home / "Documents" / "v2.txt"
    a.write_text("Mælk\nÆg\nBrød\n", "utf-8")
    b.write_text("Mælk 2 l\nÆg\nBrød\nKaffe\n", "utf-8")
    out = call(ctx, "compare_files", first=str(a), second=str(b))
    assert out["detail"]["added"] == 1 and out["detail"]["changed"] == 1
    assert 'Changed "Mælk" to "Mælk 2 l"' in out["result"] and "Kaffe" in out["result"]


def test_compare_word_versions_and_identical(ctx, home):
    for name, para in (("A", "Mødet er mandag"), ("B", "Mødet er tirsdag")):
        call(ctx, "create_document", format="docx", title="Referat", name=name, sections=[{"paragraphs": [para]}])
    out = call(ctx, "compare_files", first="A.docx", second="B.docx")
    assert "mandag" in out["result"] and "tirsdag" in out["result"]
    call(ctx, "create_document", format="docx", title="Referat", name="C", sections=[{"paragraphs": ["Mødet er mandag"]}])
    assert "same text" in call(ctx, "compare_files", first="A.docx", second="C.docx")["result"]


def test_compare_spreadsheet_cells(ctx, home):
    call(ctx, "create_document", format="xlsx", title="B1", table=[["Post", "Kr"], ["Mad", "3000"]])
    call(ctx, "create_document", format="xlsx", title="B2", table=[["Post", "Kr"], ["Mad", "3500"], ["Bus", "400"]])
    out = call(ctx, "compare_files", first="B1.xlsx", second="B2.xlsx")
    assert "B2: 3000 -> 3500" in out["result"] and out["detail"]["added"] == 2


# ---- convert ------------------------------------------------------------------------------------

def test_word_to_pdf_without_office(ctx, home):
    call(ctx, "create_document", format="docx", title="Ansøgning", sections=[{"heading": "Om mig", "paragraphs": ["Jeg hedder Mikkel"]}])
    out = call(ctx, "convert_file", path="Ansøgning.docx", to="pdf")
    p = grace_dir(home) / "Ansøgning.pdf"
    assert out["ok"] and p.exists() and "built-in converter" in out["result"] and "original is kept" in out["result"]
    assert "Jeg hedder Mikkel" in call(ctx, "read_document", path=str(p))["result"]


def test_word_to_pdf_uses_office_when_there(ctx, home, monkeypatch):
    call(ctx, "create_document", format="docx", title="W", sections=[{"paragraphs": ["x"]}])

    def fake_office(src, dst):
        dst.write_bytes(b"%PDF-1.4")
        return "Word"
    monkeypatch.setattr(D, "office_to_pdf", fake_office)
    assert "made with Word" in call(ctx, "convert_file", path="W.docx", to="pdf")["result"]


def test_markdown_to_word_and_pdf(ctx, home):
    (home / "Documents" / "plan.md").write_text("# Plan\n\n## Uge 1\n\n- løb\n- svøm\n\n| Dag | Km |\n|---|---|\n| man | 5 |\n", "utf-8")
    assert call(ctx, "convert_file", path="dokumenter/plan.md", to="word")["ok"]
    t, blocks = D.docx_blocks(home / "Documents" / "plan.docx")
    assert t == "Plan" and ("ul", ["løb", "svøm"]) in blocks and ("table", [["Dag", "Km"], ["man", "5"]]) in blocks
    assert call(ctx, "convert_file", path="dokumenter/plan.md", to="pdf")["ok"]


def test_csv_to_excel_and_back(ctx, home):
    (home / "Documents" / "tal.csv").write_text("navn,tal\nA,1\nB,2,5\n", "utf-8")
    assert call(ctx, "convert_file", path="dokumenter/tal.csv", to="excel")["ok"]
    import openpyxl
    assert openpyxl.load_workbook(home / "Documents" / "tal.xlsx").active["B2"].value == 1
    out = call(ctx, "convert_file", path="dokumenter/tal.xlsx", to="csv")
    assert out["ok"] and "tal (2).csv" in out["result"]  # never over the original csv
    assert "A" in (home / "Documents" / "tal (2).csv").read_text("utf-8-sig")


def test_images(ctx, home):
    from PIL import Image
    Image.new("RGBA", (20, 10), (255, 0, 0, 0)).save(home / "Pictures" / "logo.png")
    assert call(ctx, "convert_file", path="billeder/logo.png", to="jpg")["ok"]
    assert call(ctx, "convert_file", path="billeder/logo.png", to="webp")["ok"]
    with Image.open(home / "Pictures" / "logo.jpg") as im:
        assert im.size == (20, 10) and im.getpixel((0, 0)) == (255, 255, 255)  # transparent -> white
    assert not call(ctx, "convert_file", path="billeder/logo.png", to="docx")["ok"]


def test_pdf_to_word(ctx, home):
    call(ctx, "create_document", format="pdf", title="Kvittering", sections=[{"paragraphs": ["Beløb 450 kr"]}])
    out = call(ctx, "convert_file", path="Kvittering.pdf", to="docx")
    assert out["ok"] and "text only" in out["result"]
    assert "Beløb 450 kr" in call(ctx, "read_document", path="Kvittering.docx")["result"]


def test_convert_refusals_and_dry_run(ctx, dry_ctx, home):
    (home / "Documents" / "a.txt").write_text("x")
    assert "already" in call(ctx, "convert_file", path="dokumenter/a.txt", to="txt")["result"]
    assert "don't know" in call(ctx, "convert_file", path="dokumenter/a.txt", to="mp3")["result"]
    out = call(dry_ctx, "convert_file", path="dokumenter/a.txt", to="pdf")
    assert out["result"].startswith("Would") and not (home / "Documents" / "a.pdf").exists()


# ---- search and recent ----------------------------------------------------------------------------

def test_search_inside_files(ctx, home):
    (home / "Documents" / "noter.txt").write_text("Husk at betale husleje den første", "utf-8")
    call(ctx, "create_document", format="docx", title="Budget", sections=[{"paragraphs": ["Husleje er 8500 kr"]}])
    (home / "Documents" / "node_modules").mkdir()
    (home / "Documents" / "node_modules" / "x.txt").write_text("husleje", "utf-8")
    out = call(ctx, "search_in_files", text="husleje")
    assert out["ok"] and "2 file(s)" in out["result"] and "Budget.docx" in out["result"]
    assert "node_modules" not in out["result"]
    out = call(ctx, "search_in_files", text="betale første", folder="dokumenter")  # all words, not the phrase
    assert "noter.txt" in out["result"]
    assert call(ctx, "search_in_files", text="zebra")["status"] == "not_found"


def test_recent_files_newest_first_and_by_kind(ctx, home):
    now = time.time()
    for i, name in enumerate(["gammel.pdf", "ny.pdf", "billede.png"]):
        p = home / "Downloads" / name
        p.write_bytes(b"x")
        os.utime(p, (now - 86400 * (10 - 4 * i), now - 86400 * (10 - 4 * i)))
    out = call(ctx, "recent_files", folder="overførsler")
    assert out["result"].index("billede.png") < out["result"].index("ny.pdf") < out["result"].index("gammel.pdf")
    out = call(ctx, "recent_files", kind="pdf", days=7)
    assert "ny.pdf" in out["result"] and "gammel.pdf" not in out["result"] and "billede" not in out["result"]


def test_recent_files_uses_the_index(ctx, home):
    class Index:
        entries = [(str(home / "Desktop" / "a.docx"), "a.docx", False, time.time(), ""),
                   (str(home / "Desktop" / "mappe"), "mappe", True, time.time(), "")]
    ctx.files = Index()
    out = call(ctx, "recent_files")
    assert "a.docx" in out["result"] and "mappe" not in out["result"]


# ---- a multi-step task ------------------------------------------------------------------------------

def test_shopping_list_word_then_pdf_then_read(ctx, home):
    out = call(ctx, "create_document", format="docx", title="Indkøbsliste", sections=[{"bullets": ["mælk", "æg", "kaffe"]}])
    path = out["detail"]["path"]
    out = call(ctx, "convert_file", path=path, to="pdf")
    assert out["ok"]
    text = call(ctx, "read_document", path=out["detail"]["path"])["result"]
    assert "kaffe" in text and "æg" in text
    assert call(ctx, "undo_last")["ok"] and not Path(out["detail"]["path"]).exists() and Path(path).exists()


def test_every_files_tool_registered_with_sane_risk():
    tools = {t.name: t for t in grace_tools.load().values() if t.group in ("files", "files_docs")}
    assert sum(t.core for t in tools.values()) <= 2
    destructive = {n for n, t in tools.items() if t.risk == "destructive"}
    assert destructive == {"edit_file", "replace_file", "delete_files"}
    reads = {n for n, t in tools.items() if t.risk == "read"}
    assert reads == {"read_document", "search_in_files", "recent_files", "compare_files"}
